"""Weekly Hours Excel tab — formatting and formulas."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ===========================================================================
# Style constants
# ===========================================================================
NAVY = "2C3E7B"
GOLD = "7B5E1A"
GREEN = "1B7A3D"
PURPLE = "5B2D8B"
AMBER = "F59E0B"
BACKUP_CREAM = "FEF3C7"
TOTALS_GREEN = "B6D7A8"
EMERGENCY_RED = "FECACA"
WHITE = "FFFFFF"
GRID_GRAY = "D9D9D9"
DIVIDER_BLACK = "000000"


def _fill(argb):
    return PatternFill(start_color=argb, end_color=argb, fill_type="solid")


def _thin():
    return Side(style="thin", color=GRID_GRAY)


def _divider():
    return Side(style="medium", color=DIVIDER_BLACK)


def _section_color(col, weeks_end):
    if col <= 5:
        return NAVY
    if col <= weeks_end:
        return GOLD
    if col <= weeks_end + 2:
        return GREEN
    return PURPLE


# ===========================================================================
# Main export
# ===========================================================================


def write_weekly_sheet(wb, data, title="Weekly Hours"):
    """Add a formatted 'Weekly Hours' tab. Returns (work_refs, tag_refs)
    mapping id(emp) -> cell reference strings for Invoice live links."""
    employees = []
    for label, lst in [("Q1", data["q1"]), ("Q2", data["q2"]), ("Q3", data["q3"])]:
        if lst:
            employees.append({"fullName": label, "isSeparator": True})
            employees.extend(lst)

    first = next((e for e in employees if not e.get("isSeparator")), None)
    if not first:
        print("No weekly data")
        return {}, {}

    week_count = len(first["weekHours"])
    total_cols = 5 + week_count + 4
    abs_col = 6 + week_count
    comments_col = abs_col + 2
    weeks_end = 5 + week_count

    ws = wb.create_sheet(title)
    widths = [25, 10, 10, 15, 15] + [10] * week_count + [10, 10, 40, 35]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    header_font = Font(bold=True, color=WHITE, size=10)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    divider_cols = {5, weeks_end, abs_col + 1, comments_col, total_cols}

    def make_border(col, top_dark=False, bottom_dark=False):
        return Border(
            left=_divider() if (col == 1 or (col - 1) in divider_cols) else _thin(),
            right=_divider() if col in divider_cols else _thin(),
            top=_divider() if top_dark else _thin(),
            bottom=_divider() if bottom_dark else _thin(),
        )

    # Row 1 — Week headers
    for c in range(1, total_cols + 1):
        cell = ws.cell(1, c)
        cell.fill = _fill(_section_color(c, weeks_end))
        cell.font = header_font
        cell.alignment = center
        cell.border = make_border(c, top_dark=True)
    for i in range(week_count):
        ws.cell(1, 6 + i).value = f"Week {i + 1}"

    # Row 2 — Column headers
    headers = (
        ["Employee", "Role", "Wave", "Product", "Pilar"]
        + ["Hrs"] * week_count
        + ["Abs Hrs", "Work Hrs", "Comments", "Tag"]
    )
    for i, h in enumerate(headers, 1):
        cell = ws.cell(2, i, h)
        cell.fill = _fill(_section_color(i, weeks_end))
        cell.font = header_font
        cell.alignment = center
        cell.border = make_border(i, bottom_dark=True)

    # Data rows
    work_refs, tag_refs = {}, {}
    work_col = get_column_letter(abs_col + 1)
    tag_col = get_column_letter(comments_col + 1)
    row_idx = 3

    for emp in employees:
        if emp.get("isSeparator"):
            ws.cell(row_idx, 1, emp["fullName"])
            for c in range(1, total_cols + 1):
                cell = ws.cell(row_idx, c)
                cell.fill = _fill(AMBER)
                cell.font = Font(bold=True, color=WHITE)
                cell.alignment = wrap
                cell.border = make_border(c)
            row_idx += 1
            continue

        is_backup = emp.get("isBackup", False)
        is_emergency = not is_backup and emp.get("absHrs", 0) > 80

        ws.cell(row_idx, 1, emp["fullName"])
        ws.cell(row_idx, 2, emp["role"])
        ws.cell(row_idx, 3, emp.get("wave", ""))
        ws.cell(row_idx, 4, emp["product"])
        ws.cell(row_idx, 5, emp["pilar"])

        hrs_cells = []
        for wi, hrs in enumerate(emp["weekHours"]):
            col = 6 + wi
            ws.cell(row_idx, col, hrs).alignment = center
            hrs_cells.append(f"{get_column_letter(col)}{row_idx}")

        ws.cell(
            row_idx, abs_col, "N/A" if emp["absHrs"] == -1 else emp["absHrs"]
        ).alignment = center
        ws.cell(
            row_idx,
            abs_col + 1,
            f"=SUM({','.join(hrs_cells)})" if hrs_cells else emp["workHrs"],
        ).alignment = center
        work_refs[id(emp)] = f"'{title}'!{work_col}{row_idx}"
        ws.cell(row_idx, comments_col, emp["comments"]).alignment = wrap
        ws.cell(row_idx, comments_col + 1, emp["tag"]).alignment = center
        tag_refs[id(emp)] = f"'{title}'!{tag_col}{row_idx}"

        for c in range(1, total_cols + 1):
            cell = ws.cell(row_idx, c)
            cell.border = make_border(c)
            if not cell.alignment.horizontal:
                cell.alignment = center
            if is_backup:
                cell.fill = _fill(BACKUP_CREAM)
            elif is_emergency and (
                not cell.fill.start_color.index
                or cell.fill.start_color.index == "00000000"
            ):
                cell.fill = _fill(EMERGENCY_RED)
        row_idx += 1

    # Totals row
    data_start, data_end = 3, row_idx - 1
    ws.cell(row_idx, 1, "Totals").font = Font(bold=True)
    for wi in range(week_count):
        col = 6 + wi
        letter = get_column_letter(col)
        ws.cell(row_idx, col, f"=SUM({letter}{data_start}:{letter}{data_end})")
        ws.cell(row_idx, col).font = Font(bold=True)
        ws.cell(row_idx, col).alignment = center
    for col in (abs_col, abs_col + 1):
        letter = get_column_letter(col)
        ws.cell(row_idx, col, f"=SUM({letter}{data_start}:{letter}{data_end})")
        ws.cell(row_idx, col).font = Font(bold=True)
        ws.cell(row_idx, col).alignment = center
    total_backup = sum(
        e["workHrs"] for e in employees if e.get("isBackup") and e.get("workHrs", 0) > 0
    )
    ws.cell(row_idx, comments_col, f"Total Backup Hours: {total_backup}").font = Font(
        bold=True
    )
    ws.cell(row_idx, comments_col).alignment = center
    for c in range(1, total_cols + 1):
        cell = ws.cell(row_idx, c)
        cell.fill = _fill(TOTALS_GREEN)
        cell.border = make_border(c, top_dark=True, bottom_dark=True)

    ws.freeze_panes = "B3"
    return work_refs, tag_refs
